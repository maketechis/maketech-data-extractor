from io import BytesIO
from sqlalchemy import select
from sqlalchemy.orm import Session
from openpyxl import Workbook
from openpyxl.styles import Font
from app.enrichment.needs import detect_missing_fields
from app.models.core import CampaignLocation, District, Entity, EntityContact, EntitySource, EntityType, EntityWebsite, State

def _sheet(ws,headers):
    ws.append(headers)
    for cell in ws[1]: cell.font=Font(bold=True)
    ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions

def build_excel(db: Session, *, district_id: int|None=None) -> bytes:
    stmt=select(Entity).order_by(Entity.id)
    if district_id: stmt=stmt.where(Entity.district_id==district_id)
    entities=db.scalars(stmt).all()
    wb=Workbook(); ws=wb.active; ws.title="Entities"
    _sheet(ws,["ID","Entity Type","Name","State","District","Address","PIN","Phone","Email","Website","Verification Status"])
    contacts=wb.create_sheet("Contacts"); _sheet(contacts,["Entity ID","Type","Value","Confidence","Verified","Source URL"])
    sources=wb.create_sheet("Sources"); _sheet(sources,["Entity ID","Source Type","Source Record ID","Source URL","Collected At"])
    review=wb.create_sheet("Needs Review"); _sheet(review,["Entity ID","Name","Missing Fields","Website Status"])
    for e in entities:
        et=db.get(EntityType,e.entity_type_id); state=db.get(State,e.state_id) if e.state_id else None; district=db.get(District,e.district_id) if e.district_id else None
        cs=db.scalars(select(EntityContact).where(EntityContact.entity_id==e.id)).all()
        sites=db.scalars(select(EntityWebsite).where(EntityWebsite.entity_id==e.id).order_by(EntityWebsite.confidence.desc())).all()
        phones="; ".join(x.value for x in cs if x.contact_type=="phone"); emails="; ".join(x.value for x in cs if x.contact_type=="email")
        best=sites[0] if sites else None
        ws.append([e.id,et.name,e.name,state.name if state else "",district.name if district else "",e.address or "",e.pin or "",phones,emails,best.url if best else "",best.verification_status if best else ""])
        for c in cs: contacts.append([e.id,c.contact_type,c.value,c.confidence,c.verified,c.source_url or ""])
        for s in db.scalars(select(EntitySource).where(EntitySource.entity_id==e.id)).all(): sources.append([e.id,s.source_type,s.source_record_id or "",s.source_url,s.collected_at.isoformat() if s.collected_at else ""])
        needs=detect_missing_fields(db,e)
        if needs or (best and best.verification_status in ("candidate","needs_review")):
            review.append([e.id,e.name,", ".join(x.field for x in needs),best.verification_status if best else "no website"])
    progress=wb.create_sheet("District Progress"); _sheet(progress,["Campaign ID","District ID","Status","Attempts","Records Found","Records Saved","Last Error"])
    for row in db.scalars(select(CampaignLocation).order_by(CampaignLocation.campaign_id,CampaignLocation.id)).all():
        progress.append([row.campaign_id,row.district_id,row.status.value,row.attempts,row.records_found,row.records_saved,row.last_error or ""])
    stats=wb.create_sheet("Statistics"); _sheet(stats,["Metric","Value"])
    stats.append(["Entities",len(entities)]); stats.append(["With Phone",sum(1 for e in entities if db.scalar(select(EntityContact.id).where(EntityContact.entity_id==e.id,EntityContact.contact_type=="phone")))])
    stats.append(["With Email",sum(1 for e in entities if db.scalar(select(EntityContact.id).where(EntityContact.entity_id==e.id,EntityContact.contact_type=="email")))])
    stats.append(["With Website",sum(1 for e in entities if db.scalar(select(EntityWebsite.id).where(EntityWebsite.entity_id==e.id)))])
    for sheet in wb.worksheets:
        for col in sheet.columns:
            width=min(max((len(str(c.value or "")) for c in col),default=8)+2,60)
            sheet.column_dimensions[col[0].column_letter].width=width
    buf=BytesIO(); wb.save(buf); return buf.getvalue()
