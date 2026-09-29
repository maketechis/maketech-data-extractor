from io import BytesIO
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from app.database import Base
from app.exporters.excel import build_excel
from app.models.core import Country, District, Entity, EntityContact, EntityType, State

def test_excel_contains_core_sheets_and_entity():
    engine=create_engine("sqlite:///:memory:"); Base.metadata.create_all(engine)
    with Session(engine) as db:
        c=Country(name="India",iso_code="IN"); t=EntityType(name="Bookseller",slug="bookseller"); db.add_all([c,t]); db.flush()
        s=State(country_id=c.id,name="Bihar"); db.add(s); db.flush(); d=District(state_id=s.id,name="Siwan",normalized_name="siwan"); db.add(d); db.flush()
        e=Entity(entity_type_id=t.id,name="ABC Books",normalized_name="abc books",country_id=c.id,state_id=s.id,district_id=d.id); db.add(e); db.flush()
        db.add(EntityContact(entity_id=e.id,contact_type="phone",value="9876543210",confidence=1.0,verified=True)); db.commit()
        wb=load_workbook(BytesIO(build_excel(db,district_id=d.id)))
        assert {"Entities","Contacts","Sources","Needs Review","District Progress","Statistics"}.issubset(set(wb.sheetnames))
        assert wb["Entities"]["C2"].value=="ABC Books"
